"""Reproducible, synthetic test document. Requires python-docx only for generation."""

from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

root = Path(__file__).resolve().parents[1]
doc = Document()
doc.add_heading("A look beneath the page", 0)
doc.add_paragraph("Hello from Xray DOCX. This sentence lives inside a w:t text node.")
p = doc.add_paragraph("A paragraph can contain ")
p.add_run("bold text").bold = True
p.add_run(", italic text").italic = True
p.add_run(", and multiple runs.")
doc.add_heading("Structured content", 1)
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Part"
table.cell(0, 1).text = "Purpose"
table.cell(1, 0).text = "word/document.xml"
table.cell(1, 1).text = "Main document body"
doc.sections[0].header.paragraphs[0].text = "Xray DOCX sample header"
doc.sections[0].footer.paragraphs[0].text = "Made for inspection, not mystery."
doc.add_paragraph("Unicode works too: Bengaluru · café · ಕನ್ನಡ · 日本語")
sample_path = root / "public/sample.docx"
doc.save(sample_path)

extended_properties = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
ET.register_namespace("", extended_properties)
ET.register_namespace("vt", "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes")
with ZipFile(sample_path) as source:
    entries = [(entry, source.read(entry.filename)) for entry in source.infolist()]
for index, (entry, contents) in enumerate(entries):
    if entry.filename == "docProps/app.xml":
        properties = ET.fromstring(contents)
        application = properties.find(f"{{{extended_properties}}}Application")
        if application is None:
            raise ValueError("Sample has no extended-properties Application field.")
        application.text = "Xray DOCX"
        for node in properties.iter():
            if node.text is not None and not node.text.strip():
                node.text = None
            if node.tail is not None and not node.tail.strip():
                node.tail = None
        entries[index] = (entry, ET.tostring(properties, encoding="utf-8", xml_declaration=True))
with ZipFile(sample_path, "w", ZIP_DEFLATED) as target:
    for entry, contents in entries:
        target.writestr(entry, contents)
print(sample_path)
